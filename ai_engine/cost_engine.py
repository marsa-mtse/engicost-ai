# ==========================================================
# MTSE Cost Engine - Industrial Estimation Layer
# ==========================================================

import streamlit as st
import json
import re
import os
from typing import List, Optional
try:
    from pydantic import BaseModel, Field, ValidationError
except ImportError:
    class BaseModel:
        def dict(self): return {}
        def model_dump(self): return self.dict()
    def Field(*args, **kwargs): return None
    class ValidationError(Exception): pass

try:
    from utils import t
except ImportError:
    def t(a, e): return a


def _get_secret(key: str) -> Optional[str]:
    """Safely get a secret from st.secrets or environment variables."""
    try:
        val = st.secrets.get(key)
        return val if val else None
    except Exception:
        pass
    return os.environ.get(key) or None


try:
    import google.generativeai as genai
except ImportError:
    genai = None

try:
    import groq as groq_lib
except ImportError:
    groq_lib = None


def _parse_json_list(txt):
    """Helper to robustly extract a JSON list from AI text."""
    txt = txt.replace("```json", "").replace("```", "").strip()
    # Try to find a list first
    start = txt.find("[")
    end = txt.rfind("]")
    if start != -1 and end != -1:
        return json.loads(txt[start:end+1])
    # Try to find an object
    start = txt.find("{")
    end = txt.rfind("}")
    if start != -1 and end != -1:
        return json.loads(txt[start:end+1])
    return json.loads(txt)


class BOQItemModel(BaseModel):
    item: str = Field(description="Description of the item or work")
    unit: str = Field(description="Unit of measurement (e.g., m2, kg, ton, unit)")
    quantity: float = Field(description="Quantity required")

class BOQResultModel(BaseModel):
    items: List[BOQItemModel]

    def dict(self):
        return {"items": [item.dict() for item in self.items]}


class CostEngine:
    """
    Handles BOQ extraction and complex project cost estimation.
    Tries Groq first (reliable for this user), then Gemini as fallback.
    """
    def __init__(self):
        google_key = _get_secret("GEMINI_API_KEY") or _get_secret("GOOGLE_API_KEY")
        if google_key and genai:
            genai.configure(api_key=google_key.strip())

    def normalize_boq_data(self, data):
        """Normalizes keys from AI output (handles Arabic/English/vague keys)."""
        # If input is a dict (Groq wraps lists in objects), unwrap it first
        if isinstance(data, dict):
            for key in ["items", "boq", "line_items", "data", "بنود", "المقايسة", "list"]:
                if key in data and isinstance(data[key], list):
                    data = data[key]
                    break
            else:
                # If still a dict, can't normalize
                return [{"error": f"شكل البيانات غير متوقع: {list(data.keys())}"}]

        if not isinstance(data, list):
            return [{"error": f"البيانات المُعادة ليست قائمة: {type(data)}"}]

        mapping = {
            "item": ["item", "البند", "بند", "الوصف", "description", "name", "الاسم", "وصف_البند", "item_description"],
            "unit": ["unit", "الوحدة", "وحدة", "measurement", "وحدة_القياس"],
            "quantity": ["quantity", "الكمية", "كمية", "qty", "count", "العدد", "الكميه"]
        }

        normalized = []
        for entry in data:
            if not isinstance(entry, dict):
                continue

            new_entry = {}
            for std_key, aliases in mapping.items():
                found = False
                for k in entry.keys():
                    if k.lower() in [a.lower() for a in aliases]:
                        new_entry[std_key] = entry[k]
                        found = True
                        break
                if not found:
                    new_entry[std_key] = entry.get(std_key, "" if std_key != "quantity" else 0)

            try:
                if isinstance(new_entry["quantity"], str):
                    nums = re.findall(r"[-+]?\d*\.\d+|\d+", new_entry["quantity"])
                    new_entry["quantity"] = float(nums[0]) if nums else 0.0
                else:
                    new_entry["quantity"] = float(new_entry["quantity"] or 0)
            except:
                new_entry["quantity"] = 0.0

            normalized.append(new_entry)

        if not normalized:
            return [{"error": "لم يتم العثور على بنود في الاستجابة. حاول مرة أخرى أو تحقق من البيانات."}]
            
        # Pydantic Strict Validation
        try:
            validated_obj = BOQResultModel(items=normalized)
            # Pyre-friendly check
            model_dump_fn = getattr(validated_obj, "model_dump", None)
            if model_dump_fn and callable(model_dump_fn):
                res = model_dump_fn()
                if isinstance(res, dict):
                    return res.get("items", normalized)
            
            dict_fn = getattr(validated_obj, "dict", None)
            if dict_fn and callable(dict_fn):
                res = dict_fn()
                if isinstance(res, dict):
                    return res.get("items", normalized)
            return normalized
        except Exception as e:
            return [{"error": f"فشل التحقق من صحة البيانات (Pydantic Error): {str(e)}"}]

    def _call_groq(self, prompt_text, expect_json=True):
        """Call Groq API. Returns parsed data if expect_json=True, else raw text."""
        groq_key = _get_secret("GROQ_API_KEY")
        if not groq_key:
            return None, "لا يوجد GROQ_API_KEY"
        if not groq_lib:
            return None, "مكتبة groq غير مثبتة"

        last_err = "Unknown"
        try:
            client = groq_lib.Groq(api_key=groq_key.strip())
            for model_name in ["llama-3.3-70b-versatile", "llama-3.1-8b-instant", "mixtral-8x7b-32768"]:
                try:
                    response = client.chat.completions.create(
                        model=model_name,
                        messages=[{"role": "user", "content": prompt_text}],
                        max_tokens=4096,
                        temperature=0.7 if not expect_json else 0.1,
                    )
                    raw_text = response.choices[0].message.content
                    if expect_json:
                        return _parse_json_list(raw_text), None
                    return raw_text, None
                except Exception as e:
                    last_err = str(e)
                    if "429" in last_err or "decommissioned" in last_err:
                        continue
                    break
        except Exception as e:
            return None, f"Groq error: {str(e)}"
        return None, f"Groq fails: {last_err}"

    def _call_gemini_text(self, prompt_text, expect_json=True):
        """Call Gemini API. Returns parsed data if expect_json=True, else raw text."""
        if not genai:
            return None, "google-generativeai non installed"
        google_key = _get_secret("GEMINI_API_KEY") or _get_secret("GOOGLE_API_KEY")
        if not google_key:
            return None, "لا يوجد GOOGLE_API_KEY"

        last_err = "Unknown"
        for model_name in ["gemini-2.0-flash", "gemini-1.5-flash"]:
            try:
                model = genai.GenerativeModel(model_name)
                response = model.generate_content(prompt_text)
                if expect_json:
                    return _parse_json_list(response.text), None
                return response.text, None
            except Exception as e:
                last_err = str(e)
                if "404" in last_err or "429" in last_err:
                    continue
                break
        return None, f"Gemini fails: {last_err}"

    def extract_boq_items(self, content):
        """Extracts BOQ items from pasted text — tries Groq first, then Gemini."""
        prompt = f"""You are a Professional Quantity Surveyor.
Extract all line items from this Bill of Quantities (BOQ) text.

Content:
{content[:5000]}

Return a JSON array of objects. Each object must have:
- "item": string (item/work description)
- "unit": string (unit of measurement: m2, m3, kg, etc.)  
- "quantity": number

Example output:
[
  {{"item": "Concrete C25", "unit": "m3", "quantity": 50.5}},
  {{"item": "Reinforcement steel", "unit": "ton", "quantity": 3.2}}
]

Return ONLY the JSON array, no other text.
"""
        errors = []

        # 1. Try Groq first
        result, err = self._call_groq(prompt)
        if result is not None:
            normalized = self.normalize_boq_data(result)
            if normalized and "error" not in normalized[0]:
                return normalized
            errors.append(f"Groq normalization: {normalized[0].get('error', '')}")
        else:
            errors.append(f"Groq: {err}")

        # 2. Try Gemini as fallback
        result, err = self._call_gemini_text(prompt)
        if result is not None:
            normalized = self.normalize_boq_data(result)
            if normalized and "error" not in normalized[0]:
                return normalized
            errors.append(f"Gemini normalization: {normalized[0].get('error', '')}")
        else:
            errors.append(f"Gemini: {err}")

        return [{"error": " | ".join(errors)}]

    def _file_to_text(self, file_bytes, file_type):
        """Convert uploaded file to plain text for Groq processing."""
        import io

        # --- Excel ---
        if "excel" in file_type.lower() or "spreadsheet" in file_type.lower() or "xlsx" in file_type.lower():
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    lines.append(f"=== Sheet: {sheet.title} ===")
                    for row in sheet.iter_rows(values_only=True):
                        row_str = "\t".join([str(c) if c is not None else "" for c in row])
                        if row_str.strip():
                            lines.append(row_str)
                return "\n".join(lines)
            except Exception as e:
                return None, f"خطأ في قراءة Excel: {str(e)}"

        # --- PDF ---
        if "pdf" in file_type.lower():
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
                pages = []
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        pages.append(text)
                extracted = "\n".join(pages)
                if not extracted.strip():
                    return None, "لم يتم استخراج نص من PDF. قد يكون الملف صورة ممسوحة ضوئياً."
                return extracted, None
            except Exception as e:
                return None, f"خطأ في قراءة PDF: {str(e)}"

        # --- Word (.docx) ---
        if "word" in file_type.lower() or "vnd.openxmlformats-officedocument.wordprocessingml.document" in file_type.lower() or "docx" in file_type.lower():
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                fullText = []
                for para in doc.paragraphs:
                    fullText.append(para.text)
                return '\n'.join(fullText), None
            except Exception as e:
                return None, f"خطأ في قراءة Word: {str(e)}"

        return None, f"نوع الملف غير مدعوم: {file_type}"

    def extract_boq_from_file(self, file_bytes, file_type):
        """Extracts BOQ from uploaded files by converting to text, then using Groq."""
        # Step 1: Convert file to text
        result = self._file_to_text(file_bytes, file_type)
        if isinstance(result, tuple):
            text_content, err = result
        else:
            text_content, err = result, None

        if not text_content:
            return [{"error": err or "فشل تحويل الملف إلى نص."}]

        # Step 2: Extract BOQ from text using Groq
        return self.extract_boq_items(text_content)

    @staticmethod
    @st.cache_data(ttl=3600, show_spinner=False)
    def suggest_market_prices(items):
        """Uses AI (Groq/Gemini) with live market context to suggest current prices."""
        if not items:
            return {}
            
        from ai_engine.market_api import MarketEngine
        try:
            market_data = MarketEngine.get_live_prices()
            usd_prices = market_data.get('usd', {})
            local_prices = market_data.get('local', {})
            rate = market_data.get('rate', 1.0)
            
            currency = st.session_state.get("currency", "USD")
            region = st.session_state.get("region", "Global")
            
            ctx_usd = "\n".join([f"- {k}: ${v:,.2f}" for k, v in usd_prices.items()])
            ctx_local = "\n".join([f"- {k}: {v:,.2f} {currency}" for k, v in local_prices.items()])
            market_context = f"Exchange Rate: 1 USD = {rate:.2f} {currency} in {region}\n\n{currency} PRICES:\n{ctx_local}\n\nUSD PRICES:\n{ctx_usd}"
        except ImportError:
            currency = st.session_state.get("currency", "USD")
            region = st.session_state.get("region", "Global")
            market_context = f"No live market data available for {region}. Please estimate based on standard global prices."

        # Prepare a minimal list of items to send
        items_text = "\n".join([f"{i}. {item.get('item', '')} ({item.get('unit', '')})" for i, item in enumerate(items)])
        
        prompt = f"""You are a Professional Quantity Surveyor in {region}.
Estimate a suggested unit price (in {currency}) for these items.

CURRENT MARKET CONTEXT:
{market_context}

Items:
{items_text}

Return ONLY a JSON dictionary: {{"index": price_number}}.
Use market context for accuracy.
"""
        # Try Groq first
        dummy_instance = CostEngine()
        result, _ = dummy_instance._call_groq(prompt)
        if result and isinstance(result, dict) and len(result) > 0 and "error" not in result:
             # Convert values to float
             cleaned = {}
             for k, v in result.items():
                 try:
                     cleaned[str(k)] = float(v)
                 except:
                     pass
             return cleaned

        # Fallback to Gemini
        result, _ = dummy_instance._call_gemini_text(prompt)
        if result and isinstance(result, dict) and len(result) > 0 and "error" not in result:
             cleaned = {}
             for k, v in result.items():
                 try:
                     cleaned[str(k)] = float(v)
                 except:
                     pass
             return cleaned

        return {}

    def analyze_site_image(self, image_bytes):
        """Uses Gemini Vision to analyze a satellite image of a construction site."""
        if not genai:
            return None, "google-generativeai غير مثبت"
        google_key = _get_secret("GOOGLE_API_KEY")
        if not google_key:
            return None, "لا يوجد GOOGLE_API_KEY"

        try:
            from PIL import Image
            import io
            img = Image.open(io.BytesIO(image_bytes))
            
            prompt = """
You are an expert Civil Engineer and GIS remote sensing specialist.
Analyze this satellite image of a land plot or construction site.
Provide a structured technical report identifying:
1. Ground conditions (desert, urban, vegetation, etc.)
2. Visible structures or buildings
3. Nearby infrastructure (roads, water bodies)
4. Potential challenges for construction
Reply in highly professional Arabic in standard Markdown format. Do not use code blocks.
            """
            model = genai.GenerativeModel("gemini-1.5-flash")
            response = model.generate_content([prompt, img])
            return response.text, None
        except Exception as e:
            return None, f"خطأ في تحليل الصورة: {str(e)}"

    def generate_fidic_letter(self, issue_desc, project_info=""):
        """Generates a formal FIDIC-compliant legal letter based on an issue description."""
        prompt = f"""You are an expert FIDIC Contracts Administrator and Construction Lawyer.
Write a formal, professional legal letter to the Engineer/Client regarding the following issue.
        
Issue Description:
{issue_desc}

Project Context:
{project_info}

Requirements:
1. Write the letter entirely in professional, formal English.
2. Structure it as a standard formal letter (Header, Date, Reference, Subject, Body, Conclusion, Sign-off). Leave placeholders like [Date], [Company Name] where appropriate.
3. Explicitly cite the most relevant FIDIC Red Book (1999 or 2017) clauses that apply to this claim or notification (e.g., Clause 8.4 Extension of Time, Clause 20.1 Contractor's Claims).
4. Maintain a firm but professional and collaborative tone.
5. Do not surround the output with markdown tags like ```text. Just return the pure text.
"""
        result, err = self._call_groq(prompt, expect_json=False)
        if result:
            return result
            
        result, err = self._call_gemini_text(prompt, expect_json=False)
        if result:
            return result
            
        return "Error generating letter: " + str(err)

    def calculate_cost_matrix(self, items, base_prices, overhead=0.15, waste=0.05, profit=0.20):
        """Calculates full cost matrix with overhead, waste, and profit scenarios."""
        full_estimate = []
        total_direct = 0

        for i, item in enumerate(items):
            qty = float(item.get("quantity", 0) or 0)
            base_price = float(base_prices.get(str(i), 0) or 0)

            direct_cost = qty * base_price
            total_direct += direct_cost

            full_estimate.append({
                "item": item.get("item", t("بدون اسم", "Unnamed")),
                "qty": qty,
                "unit": item.get("unit", "-"),
                "base_price": base_price,
                "direct_total": direct_cost,
                "with_waste": direct_cost * (1 + waste),
                "with_overhead": direct_cost * (1 + waste + overhead),
                "final_price": direct_cost * (1 + waste + overhead + profit)
            })

        return {
            "items": full_estimate,
            "summary": {
                "total_direct": total_direct,
                "total_with_waste": total_direct * (1 + waste),
                "total_grand": total_direct * (1 + waste + overhead + profit)
            }
        }


@st.cache_resource
def get_cost_engine():
    return CostEngine()
